from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];d=Document(); sec=d.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(22)
def h(t):d.add_heading(t,1)
def p(t):d.add_paragraph(t)
def img(n,w=3):d.add_picture(str(R/'evidence'/n),width=Inches(w))
def pair(a,b,cap):
 par=d.add_paragraph();par.alignment=1
 par.add_run().add_picture(str(R/'evidence'/a),width=Inches(2.4));par.add_run('  ');par.add_run().add_picture(str(R/'evidence'/b),width=Inches(2.4))
 p(cap)
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):
  c.text=v;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');c._tc.get_or_add_tcPr().append(sh)
  for run in c.paragraphs[0].runs:run.bold=True
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 for row in t.rows:
  for c in row.cells:
   for pa in c.paragraphs:
    for run in pa.runs:run.font.size=Pt(9)
 bd=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
 t._tbl.tblPr.append(bd)
d.add_paragraph('Báo cáo thực hành Lab F2',style='Title')
p('Bố cục giao diện Flutter Material 3 trên Android')
p('Sinh viên: Nguyễn Thị Thanh Triều    |    MSSV: 231A310126')
p('Học phần INT4211    |    Project f2_layout')
p('Ứng dụng dựng màn hình Cổng thực hành LTDD với form đăng nhập, hồ sơ sinh viên và bố cục thích ứng. Khi vùng an toàn rộng từ 700 logical pixels, form và hồ sơ được đặt cạnh nhau theo tỉ lệ 3:2; dưới ngưỡng này chúng xếp dọc trong vùng cuộn.')
h('1 Giao diện dọc')
pair('portrait.png','profile.png','Hình 1 và 2. Màn hình dọc và phần hồ sơ sau khi cuộn trên Android Emulator.')
p('Lớp CNTT – LTDD, email 231A310126@vhu.edu.vn và các chỉ số thống kê là dữ liệu minh họa theo mẫu bài; có thể sửa trong lib/widgets/profile_card.dart.')
d.add_page_break();h('2 Bố cục ngang và tương tác')
img('landscape.png',6.7);p('Hình 3. Xoay ngang: form bên trái và hồ sơ bên phải.')
img('landscape_bottom.png',6.7);p('Hình 4. Cuộn ở hướng ngang để xem các thành phần phía dưới.')
d.add_page_break();h('Tương tác đăng nhập và chế độ tối')
pair('dark.png','snackbar.png','Hình 5 và 6. Chế độ tối và SnackBar đăng nhập mô phỏng.')
d.add_page_break();h('3 Phân tích lỗi tràn và kiểm thử')
p('Trong harness tools/overflow_demo.dart, vùng cuộn được cố ý bỏ, Column chứa banner và khối nhập cao 900 logical pixels. Tổng chiều cao vượt ràng buộc của màn hình nên Flutter báo BOTTOM OVERFLOWED BY 278 PIXELS. Đây là bản minh họa độc lập; APK nộp dùng lib/main.dart có SingleChildScrollView.')
img('overflow.png',2.4)
p('Hình 7. Lỗi tràn tạo có chủ đích, chụp từ máy ảo Android.')
p('Cách xử lý: dùng SingleChildScrollView để nội dung cuộn khi bàn phím làm giảm chiều cao; dùng Expanded/Flexible hoặc xuống dòng khi một hàng quá rộng. Hàng ghi nhớ trong ứng dụng chuyển thành Column khi vùng form hẹp, còn chữ trên banner dùng FittedBox để nằm trong khối cao 150.')
table(['Kiểm tra','Kết quả'],[
('Mật khẩu, checkbox, SnackBar và chuyển sáng tối','Đạt'),('Bề rộng 320, 699, 700, 900; bàn phím giả lập cao 350','Đạt'),('flutter analyze','Không có vấn đề'),('flutter test --coverage','5/5 đạt; 125/134 dòng, 93,3%'),('Build và chạy Android','APK debug build thành công; ảnh/video từ emulator')])
p('Độ phủ trên tính từ mã Dart trong lib, gồm widget. Kiểm thử tự động dùng flutter_test; ảnh và video bổ sung kiểm chứng chạy Android thực tế. Đăng nhập chỉ hiện thông báo mô phỏng theo yêu cầu lab, chưa có xác thực máy chủ.')
d.add_page_break();h('4 Đối chiếu Android XML và Flutter')
table(['Android XML','Flutter','Ví dụ trong mã nguồn'],[
('LinearLayout vertical','Column','main.dart: bố cục hẹp và form'),('LinearLayout horizontal','Row','main.dart: bố cục rộng'),('layout_weight','Expanded(flex)','main.dart: form flex 3, hồ sơ flex 2'),('padding / margin','Padding / Container','main.dart: đệm ngang 20; stat_box.dart: đệm 12'),('match_parent','double.infinity / stretch','header_banner.dart: chiều rộng banner'),('FrameLayout / RelativeLayout','Stack + Positioned','header_banner.dart: avatar ở mép dưới'),('ScrollView','SingleChildScrollView','main.dart: bao toàn bộ màn hình'),('drawable gradient và bo góc','BoxDecoration','header_banner.dart: gradient, góc dưới 24'),('themes.xml / style','ThemeData / ColorScheme','main.dart: makeTheme và TextTheme'),('layout-land','LayoutBuilder','main.dart: ngưỡng rộng 700'),('TextInputLayout / EditText','TextField + InputDecoration','main.dart: MSSV và mật khẩu'),('CardView và hàng thông tin','Card + ListTile','profile_card.dart: 3 dòng hồ sơ')])
h('5 Hai phần nâng cao')
p('NC1: MyApp là StatefulWidget giữ trạng thái sáng/tối, khai báo theme và darkTheme từ ColorScheme.fromSeed. IconButton trên banner gọi callback để đổi ThemeMode; dữ liệu form giữ nguyên trong lần chuyển theme.')
p('NC2: HeaderBanner, ProfileCard và StatBox nằm ở ba file riêng trong lib/widgets. StatBox bỏ dấu gạch dưới đầu tên để import được giữa các file. LoginPage ghép các widget và chọn một hoặc hai cột.')
d.add_page_break();h('6 Trả lời câu hỏi ôn tập')
answers=[
('Ba quy tắc bố cục Flutter','Ràng buộc đi xuống, kích thước đi lên, cha đặt vị trí con. Ví dụ LayoutBuilder nhận bề rộng từ vùng an toàn; Row chia phần rộng cho form và hồ sơ qua Expanded; mỗi con tính chiều cao trong giới hạn được cấp, rồi Row đặt chúng cạnh nhau.'),
('Expanded và Flexible khác nhau thế nào','Expanded dùng fit tight nên con phải lấp đầy phần được chia. Flexible mặc định fit loose, cho phép con nhỏ hơn phần đó. Trong Row, các con không flex có thể nhận bề rộng không giới hạn; Expanded cấp bề rộng hữu hạn cho Divider và giúp hai đường kẻ chia đều phần còn lại quanh chữ hoặc.'),
('Khi nào dùng Stack','Khi các phần tử cần chồng lên nhau. Banner dùng Stack với Positioned để avatar đè lên mép dưới ảnh bìa; Column và Row sắp xếp tuần tự nên không thể tạo hiệu ứng này trực tiếp.'),
('Hai cách xử lý RenderFlex overflowed','Bọc nội dung dài trong SingleChildScrollView, hoặc dùng Expanded/Flexible để giới hạn và chia không gian cho con trong Row/Column. Màn hình nhập liệu phù hợp với vùng cuộn vì bàn phím làm giảm chiều cao. Không đặt Expanded theo chiều dọc bên trong vùng cuộn không giới hạn chiều cao.'),
('So sánh XML layout-land và Flutter LayoutBuilder','XML có file riêng cho hướng ngang nên dễ thiết kế từng bố cục nhưng lặp mã và cần cập nhật đồng thời. LayoutBuilder quyết định theo bề rộng thật của vùng chứa trong một cây widget, thuận tiện cho nhiều kích thước; nhược điểm là logic và cấu trúc có thể dài, nên cần tách widget. Hướng ngang chưa chắc đạt ngưỡng 700: bố cục đổi theo chiều rộng, không chỉ theo hướng thiết bị.')]
for i,(q,a) in enumerate(answers,1):
 d.add_heading(f'{i} {q}',2);p(a)
h('7 Chạy và nộp bài')
p('Mở thư mục f2_layout bằng Android Studio có plugin Flutter/Dart, cấu hình Flutter SDK rồi chạy lib/main.dart. Các lệnh: flutter pub get; flutter run; flutter analyze; flutter test --coverage; flutter build apk --debug.')
p('Nộp F2_231A310126.zip gồm lib/ và pubspec.yaml theo đề, báo cáo này và video F2_231A310126_demo.mp4 dưới 2 phút. Bản project đầy đủ và APK debug được cung cấp thêm để mở và chạy Android dễ dàng.')
for st in d.styles:
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for par in d.paragraphs:
 for el in list(par._p.iter(qn('w:pBdr'))):el.getparent().remove(el)
d.save(R/'submission/Bao_cao_Lab_F2_231A310126.docx')
